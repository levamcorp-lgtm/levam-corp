import datetime as dt
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule, FormulaRule, DataBarRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.comments import Comment
from openpyxl.chart import BarChart, DoughnutChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.utils import get_column_letter as CL

OUT = sys.argv[1]

# ---------- Paleta corporativa ----------
NAVY = "0B2545"
NAVY2 = "13315C"
TEAL = "0E9F8A"
GOLD = "F2A900"
RED = "C0392B"
LIGHT = "F3F6FA"
MID = "DCE3ED"
WHITE = "FFFFFF"
INPUT_FILL = "FFF4C2"
GREEN_SOFT = "E3F5EF"
RED_SOFT = "FBE4E1"
FONT = "Arial"

CUR = '$#,##0.00;[Red]($#,##0.00);"-"'
CUR3 = '$#,##0.000;[Red]($#,##0.000);"-"'
PCT = '0.0%;[Red]-0.0%;"-"'
INT = '#,##0;[Red](#,##0);"-"'
DATE = "dd/mm/yyyy"

thin = Side(style="thin", color=MID)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
BOTTOM_THICK = Border(bottom=Side(style="medium", color=NAVY))
TOP_DOUBLE = Border(top=Side(style="thin", color=NAVY), bottom=Side(style="double", color=NAVY))


def fill(c):
    return PatternFill("solid", start_color=c, end_color=c)


def f(size=10, bold=False, color="1A1A1A", italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


INPUT_FONT = f(color="0000FF")
LINK_FONT = f(color="008000")

wb = Workbook()


def new_sheet(name, tab, first=False):
    ws = wb.active if first else wb.create_sheet(name)
    ws.title = name
    ws.sheet_properties.tabColor = tab
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    ws.column_dimensions["A"].width = 2.5
    return ws


def banner(ws, title, subtitle, last_col):
    ws.row_dimensions[1].height = 34
    ws.row_dimensions[2].height = 20
    for r in (1, 2):
        for c in range(1, last_col + 1):
            ws.cell(r, c).fill = fill(NAVY)
    ws["B1"] = title
    ws["B1"].font = f(18, True, WHITE)
    ws["B1"].alignment = Alignment(vertical="center")
    ws["B2"] = subtitle
    ws["B2"].font = f(10, False, "C9D6E8", italic=True)
    for c in range(1, last_col + 1):
        ws.cell(3, c).fill = fill(GOLD)
    ws.row_dimensions[3].height = 4


def header_row(ws, row, headers, col0=2, color=NAVY2, height=32):
    ws.row_dimensions[row].height = height
    for i, h in enumerate(headers):
        c = ws.cell(row, col0 + i, h)
        c.font = f(9, True, WHITE)
        c.fill = fill(color)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER


def section(ws, row, col0, col1, text, color=TEAL):
    for c in range(col0, col1 + 1):
        ws.cell(row, c).fill = fill(color)
    ws.cell(row, col0, text).font = f(10, True, WHITE)
    ws.row_dimensions[row].height = 20


def style_cell(c, fmt=None, font=None, fill_c=None, align="right", bold=False):
    if fmt:
        c.number_format = fmt
    c.font = font or f(10, bold)
    if fill_c:
        c.fill = fill(fill_c)
    c.alignment = Alignment(horizontal=align, vertical="center")
    c.border = BORDER


def note(c, text):
    cm = Comment(text, "COTIMIX")
    cm.width, cm.height = 320, 120
    c.comment = cm


def name(nm, ref):
    wb.defined_names[nm] = DefinedName(nm, attr_text=ref)


# =====================================================================
# DASHBOARD (se crea primero para que sea la primera pestaña)
# =====================================================================
dash = new_sheet("DASHBOARD", GOLD, first=True)
guia = new_sheet("GUIA", "7F8C8D")
cfg = new_sheet("CONFIG", "7F8C8D")
inv = new_sheet("INVENTARIO", TEAL)
com = new_sheet("COMPRAS", TEAL)
ven = new_sheet("VENTAS", TEAL)
gas = new_sheet("GASTOS", TEAL)
wal = new_sheet("WALMART", "0071CE")
ban = new_sheet("BANCO", "0071CE")
er = new_sheet("ESTADO_RESULTADOS", NAVY)
bal = new_sheet("BALANCE", NAVY)

# =====================================================================
# CONFIG
# =====================================================================
banner(cfg, "COTIMIX · PARÁMETROS DEL SISTEMA", "Supuestos generales y listas maestras. Celdas amarillas = editables.", 12)
cfg.column_dimensions["B"].width = 40
cfg.column_dimensions["C"].width = 22
cfg.column_dimensions["D"].width = 52
cfg.column_dimensions["E"].width = 3
for col in "FGHIJ":
    cfg.column_dimensions[col].width = 28

section(cfg, 5, 2, 4, "PARÁMETROS GENERALES")
header_row(cfg, 6, ["Parámetro", "Valor", "Nota / Fuente"])
params = [
    ("Nombre del negocio", "COTIMIX", None, None, "Dato del propietario."),
    ("Canal de venta", "Walmart Marketplace", None, None, "Dato del propietario."),
    ("Moneda", "USD", None, None, "Todas las cifras en dólares estadounidenses."),
    ("Fecha de apertura del sistema", dt.date(2026, 10, 28), DATE, "CFG_FECHA", "Fecha de corte de los saldos iniciales."),
    ("Año del Estado de Resultados", 2026, "0", "CFG_ANIO", "Cambie el año para ver otro ejercicio en ESTADO_RESULTADOS."),
    ("Capital inicial invertido", 5000, CUR, "CFG_CAPITAL", "Dato del propietario. Se registra automáticamente en BANCO como aporte."),
    ("Reserva Walmart al inicio", 1190.60, CUR, "CFG_RESERVA", "Dato del propietario: dinero de COTIMIX retenido por Walmart, aún no disponible."),
    ("Resultados acumulados de apertura", 1190.60, CUR, "CFG_RES_APERTURA",
     "SUPUESTO: la reserva de Walmart proviene de ventas anteriores a la apertura del sistema, "
     "por lo que se reconoce como utilidad acumulada. Si tuvo otro origen, ajuste este valor; "
     "el control de cuadre del BALANCE le indicará cualquier diferencia."),
    ("Recargo por pago con tarjeta Citi", 0.029, PCT, "CFG_CITI",
     "Dato del propietario: 2.9% que se suma al costo de toda compra o gasto pagado con Citi."),
    ("Comisión Walmart por defecto (referral fee)", 0.15, PCT, "CFG_WALMART",
     "SUPUESTO: 15% es la tarifa estándar de Walmart Marketplace para la mayoría de categorías. "
     "Verifique la tarifa de cada categoría en Seller Center y ajústela por producto en INVENTARIO."),
    ("Margen neto objetivo por producto", 0.25, PCT, "CFG_MARGEN",
     "Se usa para calcular el precio sugerido en INVENTARIO."),
]
for i, (lab, val, fmt, nm, nt) in enumerate(params):
    r = 7 + i
    style_cell(cfg.cell(r, 2, lab), align="left")
    c = cfg.cell(r, 3, val)
    style_cell(c, fmt, INPUT_FONT, INPUT_FILL, align="center")
    style_cell(cfg.cell(r, 4, nt), font=f(8, color="555555", italic=True), align="left")
    cfg.cell(r, 4).alignment = Alignment(wrap_text=True, vertical="center")
    cfg.row_dimensions[r].height = 30 if len(nt) < 90 else 54
    if nm:
        name(nm, f"CONFIG!$C${r}")
cfg["C7"].font = f(10, True, "0000FF")

# Listas maestras
LISTS = {
    "LST_PAGO_COMPRA": ("Métodos de pago · Compras", ["Efectivo/Banco", "Tarjeta Citi"]),
    "LST_PAGO_GASTO": ("Métodos de pago · Gastos", ["Efectivo/Banco", "Tarjeta Citi", "Descontado por Walmart"]),
    "LST_CATEGORIAS": ("Categorías de gasto", [
        "Publicidad (Walmart Sponsored)", "Envíos y logística", "Fulfillment / WFS", "Empaque y suministros",
        "Devoluciones y reembolsos", "Software y suscripciones", "Comisiones bancarias",
        "Servicios profesionales", "Impuestos y licencias", "Otros gastos"]),
    "LST_MOV_WALMART": ("Tipos de movimiento · Walmart", ["Retención a reserva", "Liberación de reserva", "Depósito a banco"]),
    "LST_MOV_BANCO": ("Tipos de movimiento · Banco", ["Aporte de capital", "Retiro del propietario", "Pago tarjeta Citi", "Otro ingreso"]),
}
section(cfg, 5, 6, 10, "LISTAS MAESTRAS (alimentan los menús desplegables)")
for j, (nm, (title, items)) in enumerate(LISTS.items()):
    col = 6 + j
    header_row(cfg, 6, [title], col0=col)
    for i, it in enumerate(items):
        style_cell(cfg.cell(7 + i, col, it), align="left", fill_c=LIGHT)
    L = CL(col)
    name(nm, f"CONFIG!${L}$7:${L}${6 + len(items)}")
CATS = LISTS["LST_CATEGORIAS"][1]
cfg.freeze_panes = "A7"

# =====================================================================
# COMPRAS  (fuente del inventario y su costo real)
# =====================================================================
C_FIRST, C_LAST = 6, 305
banner(com, "COTIMIX · REGISTRO DE COMPRAS DE INVENTARIO",
       "Cada compra suma unidades y costo al inventario. Si se paga con Citi, el 2.9% se integra al costo del producto.", 15)
heads = ["Fecha", "ID Producto", "Producto", "Proveedor", "Cantidad", "Costo unitario\n(factura)", "Subtotal",
         "Método de pago", "Recargo\ntarjeta %", "Recargo\ntarjeta $", "COSTO REAL\nTOTAL", "Costo real\npor unidad", "Notas"]
header_row(com, 5, heads)
widths = [12, 12, 24, 18, 10, 14, 14, 16, 11, 12, 15, 14, 34]
for i, w in enumerate(widths):
    com.column_dimensions[CL(2 + i)].width = w
com["B4"] = "Columnas amarillas = captura · blancas = automáticas"
com["B4"].font = f(8, color="555555", italic=True)

initial_buys = [
    (dt.date(2026, 10, 28), "P001", "Proveedor cafeteras", 23, 90, "Efectivo/Banco", "Compra inicial pagada en efectivo, sin recargo."),
    (dt.date(2026, 10, 28), "P002", "Proveedor Logitech", 40, 42, "Tarjeta Citi", "Compra inicial con Citi: 2.9% integrado al costo ($48.72)."),
]
for r in range(C_FIRST, C_LAST + 1):
    for col in (2, 3, 5, 6, 7, 9, 14):
        c = com.cell(r, col)
        style_cell(c, font=INPUT_FONT, fill_c=INPUT_FILL, align="center" if col in (2, 3, 9) else "right")
    com.cell(r, 2).number_format = DATE
    com.cell(r, 6).number_format = INT
    com.cell(r, 7).number_format = CUR
    com.cell(r, 5).alignment = Alignment(horizontal="left")
    com.cell(r, 14).alignment = Alignment(horizontal="left")
    com.cell(r, 4, f'=IF($C{r}="","",IFERROR(INDEX(INVENTARIO!$C$7:$C$26,MATCH($C{r},INVENTARIO!$B$7:$B$26,0)),"⚠ ID no existe"))')
    com.cell(r, 8, f'=IF($C{r}="","",$F{r}*$G{r})')
    com.cell(r, 10, f'=IF($C{r}="","",IF($I{r}="Tarjeta Citi",CFG_CITI,0))')
    com.cell(r, 11, f'=IF($C{r}="","",$H{r}*$J{r})')
    com.cell(r, 12, f'=IF($C{r}="","",$H{r}+$K{r})')
    com.cell(r, 13, f'=IF(OR($C{r}="",N($F{r})=0),"",$L{r}/$F{r})')
    for col, fmt in ((4, None), (8, CUR), (10, PCT), (11, CUR), (12, CUR), (13, CUR3)):
        style_cell(com.cell(r, col), fmt, bold=(col == 12), align="left" if col == 4 else "right")
for i, (d, pid, prov, q, cu, met, nt) in enumerate(initial_buys):
    r = C_FIRST + i
    com.cell(r, 2, d); com.cell(r, 3, pid); com.cell(r, 5, prov); com.cell(r, 6, q)
    com.cell(r, 7, cu); com.cell(r, 9, met); com.cell(r, 14, nt)
note(com["G7"], "Dato del propietario: Mouse Logitech a $42 c/u con Citi. Costo real = $42 × 1.029 = $43.218.")
tr = C_LAST + 1
com.cell(tr, 2, "TOTAL")
com.cell(tr, 6, f"=SUM(F{C_FIRST}:F{C_LAST})")
for col in (8, 11, 12):
    com.cell(tr, col, f"=SUM({CL(col)}{C_FIRST}:{CL(col)}{C_LAST})")
for col in range(2, 15):
    style_cell(com.cell(tr, col), CUR if col in (8, 11, 12) else INT, f(10, True, WHITE), NAVY)
dv = DataValidation(type="list", formula1="=INVENTARIO!$B$7:$B$26", allow_blank=True)
dv.error, dv.errorTitle = "Seleccione un ID registrado en INVENTARIO.", "ID inválido"
com.add_data_validation(dv); dv.add(f"C{C_FIRST}:C{C_LAST}")
dv = DataValidation(type="list", formula1="=LST_PAGO_COMPRA", allow_blank=True)
com.add_data_validation(dv); dv.add(f"I{C_FIRST}:I{C_LAST}")
com.freeze_panes = "D6"
com.auto_filter.ref = f"B5:N{C_LAST}"

# =====================================================================
# INVENTARIO (catálogo, costo real, precios y rentabilidad unitaria)
# =====================================================================
I_FIRST, I_LAST = 7, 26
banner(inv, "COTIMIX · INVENTARIO Y RENTABILIDAD POR PRODUCTO",
       "Stock y costo real se calculan desde COMPRAS y VENTAS. Capture el precio de venta en la columna amarilla.", 24)
section(inv, 5, 2, 9, "CATÁLOGO Y EXISTENCIAS", NAVY2)
section(inv, 5, 10, 13, "COSTO REAL (incluye recargo Citi)", TEAL)
section(inv, 5, 14, 24, "PRECIO Y RENTABILIDAD UNITARIA", "8E44AD")
heads = ["ID", "Producto", "Categoría", "SKU Walmart", "Unidades\ncompradas", "Unidades\nvendidas", "STOCK\nACTUAL",
         "Stock\nmínimo", "Estado", "Costo real\ntotal", "Costo real\nunitario", "VALOR\nINVENTARIO",
         "PRECIO DE\nVENTA", "Comisión\nWalmart %", "Comisión\nWalmart $", "Envío /\nfulfillment\npor unidad",
         "UTILIDAD\nPOR UNIDAD", "Margen\nneto %", "ROI por\nunidad %", "Precio\nmínimo\n(break-even)",
         "Precio\nsugerido\n(margen obj.)", "Utilidad\npotencial\ndel stock"]
header_row(inv, 6, heads, height=48)
widths = [8, 22, 16, 14, 11, 11, 11, 10, 13, 13, 13, 14, 13, 11, 12, 12, 13, 10, 10, 12, 12, 14]
for i, w in enumerate(widths):
    inv.column_dimensions[CL(2 + i)].width = w
products = [("P001", "Cafeteras", "Hogar / Cocina", 5), ("P002", "Mouse Logitech", "Electrónica", 8)]
CR = f"COMPRAS!$C${C_FIRST}:$C${C_LAST}"
V_FIRST, V_LAST = 6, 505
for r in range(I_FIRST, I_LAST + 1):
    for col in (2, 3, 4, 5, 9, 14, 15, 17):
        style_cell(inv.cell(r, col), font=INPUT_FONT, fill_c=INPUT_FILL, align="left" if col in (3, 4, 5) else "center")
    inv.cell(r, 9).number_format = INT
    inv.cell(r, 14).number_format = CUR
    inv.cell(r, 15).number_format = PCT
    inv.cell(r, 17).number_format = CUR
    g = f'$B{r}=""'
    inv.cell(r, 6, f'=IF({g},"",SUMIFS(COMPRAS!$F${C_FIRST}:$F${C_LAST},{CR},$B{r}))')
    inv.cell(r, 7, f'=IF({g},"",SUMIFS(VENTAS!$F${V_FIRST}:$F${V_LAST},VENTAS!$D${V_FIRST}:$D${V_LAST},$B{r}))')
    inv.cell(r, 8, f'=IF({g},"",F{r}-G{r})')
    inv.cell(r, 10, f'=IF({g},"",IF(H{r}<=0,"AGOTADO",IF(H{r}<=N(I{r}),"REORDENAR","OK")))')
    inv.cell(r, 11, f'=IF({g},"",SUMIFS(COMPRAS!$L${C_FIRST}:$L${C_LAST},{CR},$B{r}))')
    inv.cell(r, 12, f'=IF({g},"",IF(F{r}>0,K{r}/F{r},0))')
    inv.cell(r, 13, f'=IF({g},"",H{r}*L{r})')
    gp = f'OR($B{r}="",$N{r}="")'
    inv.cell(r, 16, f'=IF({gp},"",N{r}*O{r})')
    inv.cell(r, 18, f'=IF({gp},"",N{r}-P{r}-Q{r}-L{r})')
    inv.cell(r, 19, f'=IF({gp},"",IF(N{r}>0,R{r}/N{r},0))')
    inv.cell(r, 20, f'=IF({gp},"",IF(L{r}>0,R{r}/L{r},0))')
    inv.cell(r, 21, f'=IF({g},"",IF(O{r}<1,(L{r}+Q{r})/(1-O{r}),0))')
    inv.cell(r, 22, f'=IF({g},"",IF(O{r}+CFG_MARGEN<1,(L{r}+Q{r})/(1-O{r}-CFG_MARGEN),0))')
    inv.cell(r, 23, f'=IF({gp},"",R{r}*H{r})')
    for col, fmt in ((6, INT), (7, INT), (8, INT), (10, None), (11, CUR), (12, CUR3), (13, CUR), (16, CUR),
                     (18, CUR), (19, PCT), (20, PCT), (21, CUR), (22, CUR), (23, CUR)):
        style_cell(inv.cell(r, col), fmt, bold=col in (8, 13, 18), align="center" if col == 10 else "right")
    if r % 2 == 0:
        for col in (6, 7, 8, 11, 12, 13, 16, 18, 19, 20, 21, 22, 23):
            inv.cell(r, col).fill = fill(LIGHT)
for i, (pid, nm_, cat, smin) in enumerate(products):
    r = I_FIRST + i
    inv.cell(r, 2, pid); inv.cell(r, 3, nm_); inv.cell(r, 4, cat); inv.cell(r, 9, smin)
for r in range(I_FIRST, I_LAST + 1):
    inv.cell(r, 15).value = 0.15 if r < I_FIRST + len(products) else None
    inv.cell(r, 17).value = 0 if r < I_FIRST + len(products) else None
note(inv["N7"], "PRECIO DE VENTA: aún no definido. Escríbalo aquí y toda la rentabilidad se recalcula automáticamente.")
note(inv["O7"], "SUPUESTO: 15% (tarifa estándar Walmart). Verifique la categoría exacta en Seller Center.")
note(inv["Q7"], "Costo de envío o WFS por unidad que absorbe COTIMIX. Déjelo en 0 si el cliente lo paga.")
note(inv["I7"], "Supuesto inicial editable: unidades mínimas antes de reordenar.")
tr = I_LAST + 1
inv.cell(tr, 2, "TOTAL")
for col in (6, 7, 8, 11, 13, 23):
    inv.cell(tr, col, f"=SUM({CL(col)}{I_FIRST}:{CL(col)}{I_LAST})")
for col in range(2, 24):
    style_cell(inv.cell(tr, col), CUR if col in (11, 13, 23) else INT, f(10, True, WHITE), NAVY)
inv.row_dimensions[tr].height = 22
# Formato condicional
rng = f"J{I_FIRST}:J{I_LAST}"
inv.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"OK"'], fill=fill(GREEN_SOFT), font=f(10, True, "1E8449")))
inv.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"REORDENAR"'], fill=fill("FEF5E7"), font=f(10, True, "B9770E")))
inv.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"AGOTADO"'], fill=fill(RED_SOFT), font=f(10, True, RED)))
inv.conditional_formatting.add(f"M{I_FIRST}:M{I_LAST}", DataBarRule(start_type="num", start_value=0, end_type="max", color=TEAL))
inv.conditional_formatting.add(f"S{I_FIRST}:S{I_LAST}", FormulaRule(formula=[f'AND(ISNUMBER(S{I_FIRST}),S{I_FIRST}<CFG_MARGEN)'], font=f(10, True, RED)))
inv.conditional_formatting.add(f"S{I_FIRST}:S{I_LAST}", FormulaRule(formula=[f'AND(ISNUMBER(S{I_FIRST}),S{I_FIRST}>=CFG_MARGEN)'], font=f(10, True, "1E8449")))
inv.conditional_formatting.add(f"N{I_FIRST}:N{I_LAST}", FormulaRule(formula=[f'AND($B{I_FIRST}<>"",$N{I_FIRST}="")'], fill=fill("FFD6D6")))
inv["B29"] = "Leyenda:"
inv["B29"].font = f(9, True)
inv["C29"] = "Amarillo = dato editable · Blanco = cálculo automático · Precio en rojo claro = precio pendiente de definir · Margen en rojo = por debajo del objetivo (CONFIG)."
inv["C29"].font = f(9, color="555555", italic=True)
inv.freeze_panes = "D7"

# =====================================================================
# VENTAS
# =====================================================================
banner(ven, "COTIMIX · REGISTRO DE VENTAS WALMART MARKETPLACE",
       "Capture fecha, orden, ID, cantidad y (opcional) precio. Comisiones, costo de ventas y utilidad se calculan solos.", 19)
heads = ["Fecha", "# Orden\nWalmart", "ID Producto", "Producto", "Cantidad", "Precio unitario\n(opcional)",
         "Precio\naplicado", "VENTA\nBRUTA", "Comisión\nWalmart %", "Comisión\nWalmart $", "Envío / otros\ncargos Walmart",
         "INGRESO\nNETO", "Costo real\nunitario", "COSTO DE\nVENTAS", "UTILIDAD\nBRUTA", "Margen %", "Notas"]
header_row(ven, 5, heads)
widths = [12, 16, 11, 22, 10, 14, 12, 13, 11, 12, 14, 13, 12, 13, 13, 10, 30]
for i, w in enumerate(widths):
    ven.column_dimensions[CL(2 + i)].width = w
ven["B4"] = "Si deja vacío el precio unitario se usa el PRECIO DE VENTA del catálogo (INVENTARIO)."
ven["B4"].font = f(8, color="555555", italic=True)
IB = "INVENTARIO!$B$7:$B$26"
for r in range(V_FIRST, V_LAST + 1):
    for col in (2, 3, 4, 6, 7, 12, 18):
        style_cell(ven.cell(r, col), font=INPUT_FONT, fill_c=INPUT_FILL, align="center" if col in (3, 4) else "right")
    ven.cell(r, 2).number_format = DATE
    ven.cell(r, 6).number_format = INT
    ven.cell(r, 7).number_format = CUR
    ven.cell(r, 12).number_format = CUR
    ven.cell(r, 18).alignment = Alignment(horizontal="left")
    g = f'$D{r}=""'
    ven.cell(r, 5, f'=IF({g},"",IFERROR(INDEX(INVENTARIO!$C$7:$C$26,MATCH($D{r},{IB},0)),"⚠ ID no existe"))')
    ven.cell(r, 8, f'=IF({g},"",IF($G{r}="",IFERROR(INDEX(INVENTARIO!$N$7:$N$26,MATCH($D{r},{IB},0)),""),$G{r}))')
    ven.cell(r, 9, f'=IF({g},"",N($F{r})*IF(ISNUMBER($H{r}),$H{r},0))')
    ven.cell(r, 10, f'=IF({g},"",IFERROR(IF(INDEX(INVENTARIO!$O$7:$O$26,MATCH($D{r},{IB},0))="",CFG_WALMART,INDEX(INVENTARIO!$O$7:$O$26,MATCH($D{r},{IB},0))),CFG_WALMART))')
    ven.cell(r, 11, f'=IF({g},"",$I{r}*$J{r})')
    ven.cell(r, 13, f'=IF({g},"",$I{r}-$K{r}-N($L{r}))')
    ven.cell(r, 14, f'=IF({g},"",IFERROR(INDEX(INVENTARIO!$L$7:$L$26,MATCH($D{r},{IB},0))*1,0))')
    ven.cell(r, 15, f'=IF({g},"",N($F{r})*$N{r})')
    ven.cell(r, 16, f'=IF({g},"",$M{r}-$O{r})')
    ven.cell(r, 17, f'=IF({g},"",IF($I{r}>0,$P{r}/$I{r},""))')
    for col, fmt in ((5, None), (8, CUR), (9, CUR), (10, PCT), (11, CUR), (13, CUR), (14, CUR3), (15, CUR), (16, CUR), (17, PCT)):
        style_cell(ven.cell(r, col), fmt, bold=col in (9, 13, 16), align="left" if col == 5 else "right")
tr = V_LAST + 1
ven.cell(tr, 2, "TOTAL")
for col in (6, 9, 11, 12, 13, 15, 16):
    ven.cell(tr, col, f"=SUM({CL(col)}{V_FIRST}:{CL(col)}{V_LAST})")
ven.cell(tr, 17, f'=IF(I{tr}>0,P{tr}/I{tr},0)')
for col in range(2, 19):
    style_cell(ven.cell(tr, col), PCT if col == 17 else (INT if col == 6 else CUR), f(10, True, WHITE), NAVY)
dv = DataValidation(type="list", formula1="=INVENTARIO!$B$7:$B$26", allow_blank=True)
ven.add_data_validation(dv); dv.add(f"D{V_FIRST}:D{V_LAST}")
dv = DataValidation(type="whole", operator="greaterThan", formula1="0", allow_blank=True)
dv.error = "La cantidad debe ser un número entero mayor que 0."
ven.add_data_validation(dv); dv.add(f"F{V_FIRST}:F{V_LAST}")
dv = DataValidation(type="date", operator="greaterThan", formula1="DATE(2020,1,1)", allow_blank=True)
dv.error = "Capture una fecha válida (dd/mm/aaaa)."
ven.add_data_validation(dv); dv.add(f"B{V_FIRST}:B{V_LAST}")
ven.conditional_formatting.add(f"P{V_FIRST}:P{V_LAST}", CellIsRule(operator="lessThan", formula=["0"], fill=fill(RED_SOFT)))
ven.conditional_formatting.add(f"I{V_FIRST}:I{V_LAST}", FormulaRule(formula=[f'AND($D{V_FIRST}<>"",N($I{V_FIRST})=0)'], fill=fill("FFD6D6")))
ven.freeze_panes = "E6"
ven.auto_filter.ref = f"B5:R{V_LAST}"

# =====================================================================
# GASTOS
# =====================================================================
G_FIRST, G_LAST = 6, 305
banner(gas, "COTIMIX · REGISTRO DE GASTOS OPERATIVOS",
       "Publicidad, envíos, empaques, suscripciones, etc. Pagos con Citi suman automáticamente el 2.9%.", 12)
heads = ["Fecha", "Categoría", "Descripción", "Proveedor", "Método de pago", "Monto", "Recargo\ntarjeta %",
         "Recargo\ntarjeta $", "TOTAL REAL", "Notas"]
header_row(gas, 5, heads)
widths = [12, 28, 32, 18, 22, 13, 11, 12, 14, 30]
for i, w in enumerate(widths):
    gas.column_dimensions[CL(2 + i)].width = w
gas["B4"] = "'Descontado por Walmart' = cargos que Walmart resta de su liquidación (no salen del banco)."
gas["B4"].font = f(8, color="555555", italic=True)
for r in range(G_FIRST, G_LAST + 1):
    for col in (2, 3, 4, 5, 6, 7, 11):
        style_cell(gas.cell(r, col), font=INPUT_FONT, fill_c=INPUT_FILL, align="left")
    gas.cell(r, 2).number_format = DATE
    gas.cell(r, 7).number_format = CUR
    gas.cell(r, 7).alignment = Alignment(horizontal="right")
    g = f'N($G{r})=0'
    gas.cell(r, 8, f'=IF({g},"",IF($F{r}="Tarjeta Citi",CFG_CITI,0))')
    gas.cell(r, 9, f'=IF({g},"",$G{r}*$H{r})')
    gas.cell(r, 10, f'=IF({g},"",$G{r}+$I{r})')
    for col, fmt in ((8, PCT), (9, CUR), (10, CUR)):
        style_cell(gas.cell(r, col), fmt, bold=col == 10)
tr = G_LAST + 1
gas.cell(tr, 2, "TOTAL")
for col in (7, 9, 10):
    gas.cell(tr, col, f"=SUM({CL(col)}{G_FIRST}:{CL(col)}{G_LAST})")
for col in range(2, 12):
    style_cell(gas.cell(tr, col), CUR, f(10, True, WHITE), NAVY)
dv = DataValidation(type="list", formula1="=LST_CATEGORIAS", allow_blank=True)
gas.add_data_validation(dv); dv.add(f"C{G_FIRST}:C{G_LAST}")
dv = DataValidation(type="list", formula1="=LST_PAGO_GASTO", allow_blank=True)
gas.add_data_validation(dv); dv.add(f"F{G_FIRST}:F{G_LAST}")
gas.freeze_panes = "C6"
gas.auto_filter.ref = f"B5:K{G_LAST}"

# =====================================================================
# Ledger helper (WALMART / BANCO)
# =====================================================================
L_FIRST, L_LAST = 22, 321


def ledger(ws, list_name, prefill=None):
    header_row(ws, L_FIRST - 1, ["Fecha", "Tipo de movimiento", "Referencia / Descripción", "Monto (USD)", "Notas"])
    for r in range(L_FIRST, L_LAST + 1):
        for col in range(2, 7):
            style_cell(ws.cell(r, col), font=INPUT_FONT, fill_c=INPUT_FILL, align="left")
        ws.cell(r, 2).number_format = DATE
        ws.cell(r, 5).number_format = CUR
        ws.cell(r, 5).alignment = Alignment(horizontal="right")
    dv = DataValidation(type="list", formula1=f"={list_name}", allow_blank=True)
    ws.add_data_validation(dv); dv.add(f"C{L_FIRST}:C{L_LAST}")
    dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
    dv.error = "Capture montos positivos; el tipo de movimiento determina si suma o resta."
    ws.add_data_validation(dv); dv.add(f"E{L_FIRST}:E{L_LAST}")
    for w, col in zip([12, 26, 36, 16, 36], "BCDEF"):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = f"A{L_FIRST}"


def summary_line(ws, r, label, formula, kind="normal", fmt=CUR, nt=None):
    ws.cell(r, 2, label)
    c = ws.cell(r, 5, formula)
    for col in range(2, 6):
        cc = ws.cell(r, col)
        cc.font = f(10)
        cc.border = Border(bottom=Side(style="hair", color=MID))
    c.number_format = fmt
    c.alignment = Alignment(horizontal="right")
    if kind == "total":
        for col in range(2, 6):
            ws.cell(r, col).fill = fill(NAVY)
            ws.cell(r, col).font = f(11, True, WHITE)
    elif kind == "sub":
        for col in range(2, 6):
            ws.cell(r, col).fill = fill(MID)
            ws.cell(r, col).font = f(10, True, NAVY)
    if nt:
        ws.cell(r, 6, nt).font = f(8, color="666666", italic=True)


def SIF(sumcol, critcol, crit, ws="__"):
    return f'SUMIFS({ws}!${sumcol}${L_FIRST}:${sumcol}${L_LAST},{ws}!${critcol}${L_FIRST}:${critcol}${L_LAST},"{crit}")'


# =====================================================================
# WALMART (reserva + saldo por liquidar)
# =====================================================================
banner(wal, "COTIMIX · CUENTA WALMART MARKETPLACE Y RESERVA",
       "Dinero de COTIMIX que está en Walmart: separa la RESERVA (no disponible) del saldo por liquidar.", 8)
ledger(wal, "LST_MOV_WALMART")
section(wal, 5, 2, 5, "A. RESERVA WALMART (dinero de COTIMIX, NO disponible)", "0071CE")
summary_line(wal, 6, "Reserva al inicio (CONFIG)", "=CFG_RESERVA")
wal["E6"].font = f(10, color="008000")
summary_line(wal, 7, "(+) Retenciones enviadas a reserva", "=" + SIF("E", "C", "Retención a reserva", "WALMART"))
summary_line(wal, 8, "(−) Liberaciones de reserva", "=-" + SIF("E", "C", "Liberación de reserva", "WALMART"))
summary_line(wal, 9, "SALDO EN RESERVA (no disponible)", "=SUM(E6:E8)", "total")
section(wal, 11, 2, 5, "B. SALDO TOTAL A FAVOR EN WALMART", "0071CE")
summary_line(wal, 12, "Reserva al inicio (CONFIG)", "=CFG_RESERVA")
summary_line(wal, 13, "(+) Ingresos netos de ventas", f"=SUM(VENTAS!$M${V_FIRST}:$M${V_LAST})", nt="Venta bruta − comisión − cargos (hoja VENTAS)")
summary_line(wal, 14, "(−) Depósitos recibidos en banco", "=-" + SIF("E", "C", "Depósito a banco", "WALMART"))
summary_line(wal, 15, "(−) Cargos descontados por Walmart",
             f'=-SUMIFS(GASTOS!$J${G_FIRST}:$J${G_LAST},GASTOS!$F${G_FIRST}:$F${G_LAST},"Descontado por Walmart")', nt="Desde hoja GASTOS")
summary_line(wal, 16, "SALDO TOTAL A FAVOR EN WALMART", "=SUM(E12:E15)", "total")
summary_line(wal, 17, "   de los cuales: en reserva (no disponible)", "=E9", "sub")
summary_line(wal, 18, "   de los cuales: por liquidar (próximo depósito)", "=E16-E17", "sub")
wal["B20"] = "MOVIMIENTOS · capture aquí retenciones, liberaciones y depósitos que Walmart envía a su banco (montos siempre positivos)"
wal["B20"].font = f(9, True, NAVY)
wal.conditional_formatting.add("E18", CellIsRule(operator="lessThan", formula=["0"], fill=fill(RED_SOFT), font=f(10, True, RED)))
note(wal["E18"], "Si este saldo es negativo, hay depósitos registrados mayores a las ventas netas: revise que todas las ventas estén capturadas.")

# =====================================================================
# BANCO (efectivo + tarjeta Citi)
# =====================================================================
banner(ban, "COTIMIX · BANCO / EFECTIVO Y TARJETA CITI",
       "Saldo de efectivo disponible y deuda de la tarjeta Citi, calculados desde todos los registros.", 8)
ledger(ban, "LST_MOV_BANCO")
ban.cell(L_FIRST, 2, "=CFG_FECHA"); ban.cell(L_FIRST, 3, "Aporte de capital")
ban.cell(L_FIRST, 4, "Capital inicial invertido COTIMIX"); ban.cell(L_FIRST, 5, "=CFG_CAPITAL")
ban.cell(L_FIRST, 6, "Vinculado a CONFIG (no editar aquí)")
for col in (2, 5):
    ban.cell(L_FIRST, col).font = LINK_FONT
section(ban, 5, 2, 5, "A. EFECTIVO / BANCO DISPONIBLE", TEAL)
summary_line(ban, 6, "(+) Aportes de capital", "=" + SIF("E", "C", "Aporte de capital", "BANCO"))
summary_line(ban, 7, "(+) Depósitos recibidos de Walmart", "=-WALMART!E14", nt="Desde hoja WALMART")
summary_line(ban, 8, "(+) Otros ingresos", "=" + SIF("E", "C", "Otro ingreso", "BANCO"))
summary_line(ban, 9, "(−) Retiros del propietario", "=-" + SIF("E", "C", "Retiro del propietario", "BANCO"))
summary_line(ban, 10, "(−) Pagos a tarjeta Citi", "=-" + SIF("E", "C", "Pago tarjeta Citi", "BANCO"))
summary_line(ban, 11, "(−) Compras de inventario pagadas en efectivo",
             f'=-SUMIFS(COMPRAS!$L${C_FIRST}:$L${C_LAST},COMPRAS!$I${C_FIRST}:$I${C_LAST},"Efectivo/Banco")', nt="Desde hoja COMPRAS")
summary_line(ban, 12, "(−) Gastos pagados en efectivo",
             f'=-SUMIFS(GASTOS!$J${G_FIRST}:$J${G_LAST},GASTOS!$F${G_FIRST}:$F${G_LAST},"Efectivo/Banco")', nt="Desde hoja GASTOS")
summary_line(ban, 13, "SALDO DE EFECTIVO DISPONIBLE", "=SUM(E6:E12)", "total")
section(ban, 15, 2, 5, "B. TARJETA DE CRÉDITO CITI (pasivo)", RED)
summary_line(ban, 16, "(+) Compras de inventario con Citi (incl. 2.9%)",
             f'=SUMIFS(COMPRAS!$L${C_FIRST}:$L${C_LAST},COMPRAS!$I${C_FIRST}:$I${C_LAST},"Tarjeta Citi")')
summary_line(ban, 17, "(+) Gastos pagados con Citi (incl. 2.9%)",
             f'=SUMIFS(GASTOS!$J${G_FIRST}:$J${G_LAST},GASTOS!$F${G_FIRST}:$F${G_LAST},"Tarjeta Citi")')
summary_line(ban, 18, "(−) Pagos realizados a Citi", "=E10")
summary_line(ban, 19, "SALDO POR PAGAR A CITI", "=SUM(E16:E18)", "total")
for r in (15,):
    pass
ban.conditional_formatting.add("E13", CellIsRule(operator="lessThan", formula=["0"], fill=fill(RED_SOFT), font=f(11, True, RED)))
note(ban["E13"], "Si el saldo es negativo, hay egresos registrados sin fondos suficientes: revise aportes o depósitos de Walmart.")
note(ban["E19"], "SUPUESTO: la compra de mouse con Citi ($1,728.72) sigue pendiente de pago. "
     "Cuando pague la tarjeta, registre un movimiento 'Pago tarjeta Citi' abajo.")

# =====================================================================
# ESTADO DE RESULTADOS
# =====================================================================
banner(er, "COTIMIX · ESTADO DE RESULTADOS MENSUAL",
       "Se genera automáticamente desde VENTAS, GASTOS y BANCO. Cambie el año en CONFIG.", 17)
er.column_dimensions["B"].width = 38
for c in range(3, 15):
    er.column_dimensions[CL(c)].width = 11.5
er.column_dimensions["O"].width = 14
er.column_dimensions["P"].width = 3
er.column_dimensions["Q"].width = 16
er["B4"] = '="Ejercicio "&CFG_ANIO&"  ·  cifras en USD"'
er["B4"].font = f(11, True, NAVY)
MESES = ["ENE", "FEB", "MAR", "ABR", "MAY", "JUN", "JUL", "AGO", "SEP", "OCT", "NOV", "DIC"]
for m in range(12):
    c = er.cell(5, 3 + m, f"=DATE(CFG_ANIO,{m + 1},1)")
    c.number_format = DATE
    c.font = f(7, color="BBBBBB")
er.row_dimensions[5].hidden = True
header_row(er, 6, ["Concepto"] + MESES + ["TOTAL AÑO"], height=24)
header_row(er, 6, ["ACUMULADO\nHISTÓRICO"], col0=17, color=GOLD, height=30)
er["Q6"].font = f(9, True, NAVY)

VD = f"VENTAS!$B${V_FIRST}:$B${V_LAST}"


def per(col):
    return f'{VD},">="&{col}$5,{VD},"<"&DATE(YEAR({col}$5),MONTH({col}$5)+1,1)'


def gper(col):
    d = f"GASTOS!$B${G_FIRST}:$B${G_LAST}"
    return f'{d},">="&{col}$5,{d},"<"&DATE(YEAR({col}$5),MONTH({col}$5)+1,1)'


def bper(col):
    d = f"BANCO!$B${L_FIRST}:$B${L_LAST}"
    return f'{d},">="&{col}$5,{d},"<"&DATE(YEAR({col}$5),MONTH({col}$5)+1,1)'


def vs(colv):
    return f"VENTAS!${colv}${V_FIRST}:${colv}${V_LAST}"


rows = {}
r = 7
layout = [
    ("sec", "INGRESOS"),
    ("Ventas brutas", lambda c: f"=SUMIFS({vs('I')},{per(c)})", f"=SUM({vs('I')})", "line"),
    ("(−) Comisiones Walmart", lambda c: f"=-SUMIFS({vs('K')},{per(c)})", f"=-SUM({vs('K')})", "line"),
    ("(−) Envíos y cargos Walmart", lambda c: f"=-SUMIFS({vs('L')},{per(c)})", f"=-SUM({vs('L')})", "line"),
    ("INGRESOS NETOS", "SUM3", None, "sub"),
    ("(−) Costo de ventas (costo real)", lambda c: f"=-SUMIFS({vs('O')},{per(c)})", f"=-SUM({vs('O')})", "line"),
    ("UTILIDAD BRUTA", "UB", None, "sub"),
    ("Margen bruto %", "MB", None, "pct"),
    ("sec", "GASTOS OPERATIVOS"),
]
for cat in CATS:
    layout.append((cat, ("CAT", cat), None, "line"))
layout += [
    ("TOTAL GASTOS OPERATIVOS", "TG", None, "sub"),
    ("UTILIDAD OPERATIVA", "UO", None, "sub"),
    ("(+) Otros ingresos", lambda c: f'=SUMIFS(BANCO!$E${L_FIRST}:$E${L_LAST},BANCO!$C${L_FIRST}:$C${L_LAST},"Otro ingreso",{bper(c)})',
     f'=SUMIFS(BANCO!$E${L_FIRST}:$E${L_LAST},BANCO!$C${L_FIRST}:$C${L_LAST},"Otro ingreso")', "line"),
    ("UTILIDAD NETA", "UN", None, "total"),
    ("Margen neto %", "MN", None, "pct"),
    ("Unidades vendidas", lambda c: f"=SUMIFS({vs('F')},{per(c)})", f"=SUM({vs('F')})", "units"),
]
cat_rows = []
for item in layout:
    if item[0] == "sec":
        section(er, r, 2, 17, item[1], TEAL)
        r += 1
        continue
    label, how, hist, kind = item
    rows[label] = r
    er.cell(r, 2, label)
    cols = [CL(c) for c in range(3, 15)] + ["Q"]
    for col in cols:
        ci = er[f"{col}{r}"]
        if callable(how):
            ci.value = hist if col == "Q" else how(col)
        elif isinstance(how, tuple):
            cat = how[1]
            gs = f"GASTOS!$J${G_FIRST}:$J${G_LAST},GASTOS!$C${G_FIRST}:$C${G_LAST},\"{cat}\""
            ci.value = f"=-SUMIFS({gs})" if col == "Q" else f"=-SUMIFS({gs},{gper(col)})"
        elif how == "SUM3":
            a = rows["Ventas brutas"]
            ci.value = f"=SUM({col}{a}:{col}{a + 2})"
        elif how == "UB":
            ci.value = f"={col}{rows['INGRESOS NETOS']}+{col}{rows['(−) Costo de ventas (costo real)']}"
        elif how == "MB":
            ci.value = f"=IF({col}{rows['Ventas brutas']}>0,{col}{rows['UTILIDAD BRUTA']}/{col}{rows['Ventas brutas']},0)"
        elif how == "TG":
            ci.value = f"=SUM({col}{cat_rows[0]}:{col}{cat_rows[-1]})"
        elif how == "UO":
            ci.value = f"={col}{rows['UTILIDAD BRUTA']}+{col}{rows['TOTAL GASTOS OPERATIVOS']}"
        elif how == "UN":
            ci.value = f"={col}{rows['UTILIDAD OPERATIVA']}+{col}{rows['(+) Otros ingresos']}"
        elif how == "MN":
            ci.value = f"=IF({col}{rows['Ventas brutas']}>0,{col}{rows['UTILIDAD NETA']}/{col}{rows['Ventas brutas']},0)"
    er[f"O{r}"] = (f"=SUM(C{r}:N{r})" if kind not in ("pct",) else
                   (f"=IF(O{rows['Ventas brutas']}>0,O{rows['UTILIDAD BRUTA']}/O{rows['Ventas brutas']},0)" if how == "MB"
                    else f"=IF(O{rows['Ventas brutas']}>0,O{rows['UTILIDAD NETA']}/O{rows['Ventas brutas']},0)"))
    if isinstance(how, tuple):
        cat_rows.append(r)
    fmt = PCT if kind == "pct" else (INT if kind == "units" else CUR)
    for c in range(2, 18):
        if c == 16:
            continue
        cc = er.cell(r, c)
        cc.number_format = fmt
        cc.border = Border(bottom=Side(style="hair", color=MID))
        cc.alignment = Alignment(horizontal="left" if c == 2 else "right", vertical="center",
                                 indent=1 if (c == 2 and kind == "line") else 0)
        if kind == "sub":
            cc.font = f(10, True, NAVY); cc.fill = fill(MID)
        elif kind == "total":
            cc.font = f(11, True, WHITE); cc.fill = fill(NAVY)
        elif kind == "pct":
            cc.font = f(9, False, "555555", italic=True)
        else:
            cc.font = f(10)
        if c == 15 and kind not in ("total",):
            cc.font = f(10, True, NAVY); cc.fill = fill(LIGHT if kind != "sub" else MID)
        if c == 17 and kind not in ("total",):
            cc.fill = fill("FFF8E1"); cc.font = f(10, True, NAVY)
    r += 1
ER_UN = rows["UTILIDAD NETA"]
ER_VB = rows["Ventas brutas"]
ER_IN = rows["INGRESOS NETOS"]
ER_UB = rows["UTILIDAD BRUTA"]
ER_UNI = rows["Unidades vendidas"]
er.conditional_formatting.add(f"C{ER_UN}:Q{ER_UN}", CellIsRule(operator="lessThan", formula=["0"], font=f(11, True, "FFB3A7")))
er.freeze_panes = "C7"
er.cell(r + 1, 2, "Notas: los gastos se muestran con signo negativo. El costo de ventas usa el costo real promedio "
        "(incluye el 2.9% de Citi cuando aplica). La columna ACUMULADO HISTÓRICO incluye todos los años y alimenta el BALANCE.").font = f(8, color="666666", italic=True)

# =====================================================================
# BALANCE GENERAL
# =====================================================================
banner(bal, "COTIMIX · BALANCE GENERAL", "Posición financiera actual. Activo = Pasivo + Patrimonio.", 8)
bal.column_dimensions["B"].width = 46
bal.column_dimensions["C"].width = 18
bal.column_dimensions["D"].width = 12
bal.column_dimensions["E"].width = 50
bal["B4"] = '="Al "&TEXT(TODAY(),"dd/mm/yyyy")&"  ·  cifras en USD"'
bal["B4"].font = f(11, True, NAVY)


def bline(r, label, formula, kind="line", fmt=CUR, nt=None):
    bal.cell(r, 2, label)
    c = bal.cell(r, 3, formula)
    c.number_format = fmt
    for col in (2, 3, 4):
        cc = bal.cell(r, col)
        cc.border = Border(bottom=Side(style="hair", color=MID))
        cc.font = f(10)
    bal.cell(r, 2).alignment = Alignment(indent=1 if kind == "line" else 0)
    if kind == "sub":
        for col in (2, 3, 4):
            bal.cell(r, col).fill = fill(MID); bal.cell(r, col).font = f(10, True, NAVY)
    if kind == "total":
        for col in (2, 3, 4):
            bal.cell(r, col).fill = fill(NAVY); bal.cell(r, col).font = f(11, True, WHITE)
    if nt:
        bal.cell(r, 5, nt).font = f(8, color="666666", italic=True)


section(bal, 6, 2, 4, "ACTIVO", TEAL)
bline(7, "Efectivo / Banco disponible", "=BANCO!E13", nt="Disponible de inmediato")
bline(8, "Saldo Walmart por liquidar", "=WALMART!E18", nt="Se depositará en la próxima liquidación")
bline(9, "Reserva Walmart (restringido, no disponible)", "=WALMART!E9", nt="Pertenece a COTIMIX; Walmart lo libera según sus políticas")
bline(10, "Inventario de mercancía (costo real)", f"=INVENTARIO!M{I_LAST + 1}", nt="Incluye el 2.9% de Citi capitalizado en el mouse")
bline(11, "TOTAL ACTIVO", "=SUM(C7:C10)", "total")
for rr in (7, 8, 9, 10):
    bal.cell(rr, 4, f"=IF($C$11<>0,C{rr}/$C$11,0)").number_format = PCT
    bal.cell(rr, 4).font = f(8, color="666666", italic=True)
section(bal, 13, 2, 4, "PASIVO", RED)
bline(14, "Tarjeta de crédito Citi por pagar", "=BANCO!E19", nt="Supuesto: compra de mouse aún no pagada")
bline(15, "TOTAL PASIVO", "=SUM(C14:C14)", "total")
section(bal, 17, 2, 4, "PATRIMONIO", NAVY2)
bline(18, "Capital aportado", "=BANCO!E6")
bline(19, "(−) Retiros del propietario", "=BANCO!E9")
bline(20, "Resultados acumulados de apertura", "=CFG_RES_APERTURA", nt="Ver supuesto en CONFIG (origen de la reserva Walmart)")
bline(21, "Resultado acumulado desde la apertura", f"=ESTADO_RESULTADOS!Q{ER_UN}", nt="Utilidad neta histórica (ESTADO_RESULTADOS)")
bline(22, "TOTAL PATRIMONIO", "=SUM(C18:C21)", "total")
bline(24, "TOTAL PASIVO + PATRIMONIO", "=C15+C22", "sub")
bline(25, "Control de cuadre (Activo − Pasivo − Patrimonio)", "=ROUND(C11-C24,2)")
bal["D25"] = '=IF(ABS(C25)<0.01,"✔ CUADRA","✖ REVISAR")'
bal["D25"].font = f(10, True)
bal.conditional_formatting.add("D25", CellIsRule(operator="equal", formula=['"✔ CUADRA"'], fill=fill(GREEN_SOFT), font=f(10, True, "1E8449")))
bal.conditional_formatting.add("D25", CellIsRule(operator="notEqual", formula=['"✔ CUADRA"'], fill=fill(RED_SOFT), font=f(10, True, RED)))

section(bal, 27, 2, 4, "INDICADORES FINANCIEROS", GOLD)
bline(28, "Liquidez inmediata (Efectivo ÷ Pasivo)", "=IF(C15>0,C7/C15,0)", fmt='0.00"x"', nt="> 1.0x: el efectivo cubre la deuda Citi")
bline(29, "Liquidez ampliada ((Efectivo + Walmart por liquidar) ÷ Pasivo)", "=IF(C15>0,(C7+C8)/C15,0)", fmt='0.00"x"')
bline(30, "Endeudamiento (Pasivo ÷ Activo)", "=IF(C11>0,C15/C11,0)", fmt=PCT, nt="Proporción de activos financiados con deuda")
bline(31, "Dinero restringido (Reserva ÷ Activo)", "=IF(C11>0,C9/C11,0)", fmt=PCT)
bline(32, "Peso del inventario (Inventario ÷ Activo)", "=IF(C11>0,C10/C11,0)", fmt=PCT)
bline(33, "ROI sobre capital (Utilidad acumulada ÷ Capital)", "=IF(CFG_CAPITAL>0,(C20+C21)/CFG_CAPITAL,0)", fmt=PCT,
      nt="Incluye resultados de apertura")
for rr in range(28, 34):
    bal.cell(rr, 3).font = f(10, True, NAVY)

# =====================================================================
# DASHBOARD
# =====================================================================
for c in range(1, 15):
    dash.column_dimensions[CL(c)].width = 13.5
dash.column_dimensions["A"].width = 2.5
dash.column_dimensions["N"].width = 2.5
banner(dash, "COTIMIX  ·  DASHBOARD FINANCIERO EJECUTIVO",
       "Walmart Marketplace · Sistema de control financiero · Cifras en USD", 14)
dash["B2"] = '="Walmart Marketplace  ·  Actualizado al "&TEXT(TODAY(),"dd/mm/yyyy")&"  ·  Cifras en USD"'
dash["B2"].font = f(10, False, "C9D6E8", italic=True)


def card(row, col, title, formula, sub, accent, fmt=CUR):
    c0, c1 = col, col + 2
    dash.merge_cells(start_row=row, start_column=c0, end_row=row, end_column=c1)
    dash.merge_cells(start_row=row + 1, start_column=c0, end_row=row + 1, end_column=c1)
    dash.merge_cells(start_row=row + 2, start_column=c0, end_row=row + 2, end_column=c1)
    for rr in range(row, row + 3):
        for cc in range(c0, c1 + 1):
            cell = dash.cell(rr, cc)
            cell.fill = fill(WHITE if rr > row else accent)
            left = Side(style="thick", color=accent) if cc == c0 else None
            cell.border = Border(left=left, bottom=Side(style="thin", color=MID) if rr == row + 2 else None,
                                 right=Side(style="thin", color=MID) if cc == c1 else None)
    t = dash.cell(row, c0, title)
    t.font = f(9, True, WHITE)
    t.alignment = Alignment(horizontal="center", vertical="center")
    v = dash.cell(row + 1, c0, formula)
    v.font = f(20, True, NAVY)
    v.number_format = fmt
    v.alignment = Alignment(horizontal="center", vertical="center")
    s = dash.cell(row + 2, c0, sub)
    s.font = f(8, False, "666666", italic=True)
    s.alignment = Alignment(horizontal="center", vertical="center")
    dash.row_dimensions[row].height = 20
    dash.row_dimensions[row + 1].height = 36
    dash.row_dimensions[row + 2].height = 18


section(dash, 5, 2, 13, "POSICIÓN FINANCIERA", NAVY2)
card(6, 2, "CAPITAL INVERTIDO", "=CFG_CAPITAL", "Aporte inicial del propietario", NAVY)
card(6, 5, "EFECTIVO DISPONIBLE", "=BANCO!E13", "Dinero listo para operar", TEAL)
card(6, 8, "RESERVA WALMART", "=WALMART!E9", "Propio · NO disponible aún", "0071CE")
card(6, 11, "VALOR DEL INVENTARIO", f"=INVENTARIO!M{I_LAST + 1}", '=TEXT(INVENTARIO!H' + str(I_LAST + 1) + ',"#,##0")&" unidades a costo real"', "8E44AD")
dash.row_dimensions[9].height = 8
card(10, 2, "ACTIVOS TOTALES", "=BALANCE!C11", "Efectivo + Walmart + Inventario", NAVY)
card(10, 5, "DEUDA TARJETA CITI", "=BALANCE!C15", '=IF(BALANCE!C15>0,"Pendiente de pago","Sin deuda")', RED)
card(10, 8, "PATRIMONIO NETO", "=BALANCE!C22", "Activos − Pasivos", TEAL)
card(10, 11, "WALMART POR LIQUIDAR", "=WALMART!E18", "Próximo depósito estimado", "0071CE")
dash.row_dimensions[13].height = 8
section(dash, 14, 2, 13, "DESEMPEÑO COMERCIAL (acumulado desde la apertura)", NAVY2)
card(15, 2, "VENTAS BRUTAS", f"=ESTADO_RESULTADOS!Q{ER_VB}", '=TEXT(ESTADO_RESULTADOS!Q' + str(ER_UNI) + ',"#,##0")&" unidades vendidas"', GOLD)
card(15, 5, "UTILIDAD BRUTA", f"=ESTADO_RESULTADOS!Q{ER_UB}", "Después de comisiones y costo", TEAL)
card(15, 8, "UTILIDAD NETA", f"=ESTADO_RESULTADOS!Q{ER_UN}", "Después de gastos operativos", NAVY)
card(15, 11, "MARGEN NETO", f"=IF(ESTADO_RESULTADOS!Q{ER_VB}>0,ESTADO_RESULTADOS!Q{ER_UN}/ESTADO_RESULTADOS!Q{ER_VB},0)",
     '="Objetivo: "&TEXT(CFG_MARGEN,"0%")', "8E44AD", fmt=PCT)
for rr in (11, 16):
    for cc in (2, 5, 8, 11):
        pass
dash.conditional_formatting.add("H16", CellIsRule(operator="lessThan", formula=["0"], font=f(20, True, RED)))
dash.conditional_formatting.add("E7", CellIsRule(operator="lessThan", formula=["0"], font=f(20, True, RED)))

# Alertas / salud
dash.row_dimensions[18].height = 8
section(dash, 19, 2, 7, "ALERTAS Y CONTROLES", RED)
alerts = [
    ("Cuadre del balance", '=BALANCE!D25'),
    ("Productos sin precio de venta", f'=COUNTIFS(INVENTARIO!$B$7:$B$26,"<>",INVENTARIO!$N$7:$N$26,"")&" producto(s)"'),
    ("Productos por reordenar / agotados",
     f'=COUNTIF(INVENTARIO!$J$7:$J$26,"REORDENAR")&" por reordenar · "&COUNTIF(INVENTARIO!$J$7:$J$26,"AGOTADO")&" agotado(s)"'),
    ("Margen por debajo del objetivo", f'=COUNTIFS(INVENTARIO!$S$7:$S$26,"<"&CFG_MARGEN,INVENTARIO!$N$7:$N$26,"<>")&" producto(s)"'),
    ("Liquidez inmediata (efectivo ÷ deuda)", '=IF(BALANCE!C15>0,TEXT(BALANCE!C28,"0.00")&"x"&IF(BALANCE!C28>=1,"  ✔ cubre la deuda","  ⚠ no cubre la deuda"),"Sin deuda")'),
    ("Dinero total de COTIMIX (efectivo + Walmart)", '=TEXT(BALANCE!C7+BALANCE!C8+BALANCE!C9,"$#,##0.00")'),
]
for i, (lab, fm) in enumerate(alerts):
    rr = 20 + i
    dash.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=4)
    dash.merge_cells(start_row=rr, start_column=5, end_row=rr, end_column=7)
    dash.cell(rr, 2, lab).font = f(9, True, NAVY)
    dash.cell(rr, 5, fm).font = f(9, True)
    dash.cell(rr, 5).alignment = Alignment(horizontal="right")
    for cc in range(2, 8):
        dash.cell(rr, cc).fill = fill(LIGHT if i % 2 == 0 else WHITE)
        dash.cell(rr, cc).border = Border(bottom=Side(style="hair", color=MID))
    dash.row_dimensions[rr].height = 18
dash.conditional_formatting.add("E20", CellIsRule(operator="equal", formula=['"✔ CUADRA"'], font=f(9, True, "1E8449")))
dash.conditional_formatting.add("E20", CellIsRule(operator="notEqual", formula=['"✔ CUADRA"'], font=f(9, True, RED)))
dash.conditional_formatting.add("E21", FormulaRule(formula=['LEFT(E21,1)<>"0"'], font=f(9, True, "B9770E")))

# Composición de activos (tabla para gráfico)
section(dash, 19, 9, 13, "¿DÓNDE ESTÁ EL DINERO DE COTIMIX?", TEAL)
comp = [("Efectivo / Banco", "=MAX(BALANCE!C7,0)"), ("Walmart por liquidar", "=MAX(BALANCE!C8,0)"),
        ("Reserva Walmart", "=MAX(BALANCE!C9,0)"), ("Inventario", "=MAX(BALANCE!C10,0)")]
for i, (lab, fm) in enumerate(comp):
    rr = 20 + i
    dash.merge_cells(start_row=rr, start_column=9, end_row=rr, end_column=10)
    dash.cell(rr, 9, lab).font = f(9, True, NAVY)
    dash.cell(rr, 11, fm).number_format = CUR
    dash.cell(rr, 11).font = f(9, True)
    dash.cell(rr, 12, f"=IF(SUM($K$20:$K$23)>0,K{rr}/SUM($K$20:$K$23),0)").number_format = PCT
    dash.cell(rr, 12).font = f(9, color="666666")
    for cc in range(9, 14):
        dash.cell(rr, cc).fill = fill(LIGHT if i % 2 == 0 else WHITE)
        dash.cell(rr, cc).border = Border(bottom=Side(style="hair", color=MID))
dash.conditional_formatting.add("M20:M23", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color=TEAL))
for i in range(4):
    dash.cell(20 + i, 13, f"=L{20 + i}").number_format = ";;;"

# Tabla de productos
dash.row_dimensions[26].height = 8
section(dash, 27, 2, 13, "INVENTARIO Y RENTABILIDAD POR PRODUCTO", NAVY2)
hd = ["Producto", "", "Stock", "Costo real\nunitario", "Precio\nventa", "Utilidad\nunitaria", "Margen\nneto", "Valor\ninventario",
      "Unid.\nvendidas", "Precio\nsugerido", "Estado", ""]
header_row(dash, 28, hd, col0=2, color=TEAL, height=30)
dash.merge_cells("B28:C28")
dash.merge_cells("L28:M28")
dash["L28"] = "Estado"
dash["L28"].font = f(9, True, WHITE)
dash["L28"].alignment = Alignment(horizontal="center", vertical="center")
dash["K28"] = "Precio\nsugerido"
NPROD_DASH = 8
for i in range(NPROD_DASH):
    rr = 29 + i
    ir = I_FIRST + i
    dash.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=3)
    dash.merge_cells(start_row=rr, start_column=12, end_row=rr, end_column=13)
    vals = [(2, f'=IF(INVENTARIO!B{ir}="","",INVENTARIO!C{ir})', None),
            (4, f'=IF(INVENTARIO!B{ir}="","",INVENTARIO!H{ir})', INT),
            (5, f'=IF(INVENTARIO!B{ir}="","",INVENTARIO!L{ir})', CUR3),
            (6, f'=IF(INVENTARIO!B{ir}="","",IF(INVENTARIO!N{ir}="","Definir",INVENTARIO!N{ir}))', CUR),
            (7, f'=IF(INVENTARIO!R{ir}="","",INVENTARIO!R{ir})', CUR),
            (8, f'=IF(INVENTARIO!S{ir}="","",INVENTARIO!S{ir})', PCT),
            (9, f'=IF(INVENTARIO!B{ir}="","",INVENTARIO!M{ir})', CUR),
            (10, f'=IF(INVENTARIO!B{ir}="","",INVENTARIO!G{ir})', INT),
            (11, f'=IF(INVENTARIO!B{ir}="","",INVENTARIO!V{ir})', CUR),
            (12, f'=IF(INVENTARIO!B{ir}="","",INVENTARIO!J{ir})', None)]
    for col, fm, fmt in vals:
        c = dash.cell(rr, col, fm)
        if fmt:
            c.number_format = fmt
        c.font = f(9, col in (2, 9))
        c.alignment = Alignment(horizontal="left" if col == 2 else "center", vertical="center")
    for cc in range(2, 14):
        dash.cell(rr, cc).fill = fill(LIGHT if i % 2 == 0 else WHITE)
        dash.cell(rr, cc).border = Border(bottom=Side(style="hair", color=MID))
    dash.row_dimensions[rr].height = 18
tr = 29 + NPROD_DASH
dash.merge_cells(start_row=tr, start_column=2, end_row=tr, end_column=3)
dash.cell(tr, 2, "TOTAL")
dash.cell(tr, 4, f"=INVENTARIO!H{I_LAST + 1}").number_format = INT
dash.cell(tr, 9, f"=INVENTARIO!M{I_LAST + 1}").number_format = CUR
dash.cell(tr, 10, f"=INVENTARIO!G{I_LAST + 1}").number_format = INT
for cc in range(2, 14):
    dash.cell(tr, cc).fill = fill(NAVY)
    dash.cell(tr, cc).font = f(9, True, WHITE)
    dash.cell(tr, cc).alignment = Alignment(horizontal="center" if cc > 3 else "left")
pr = f"L29:L{28 + NPROD_DASH}"
dash.conditional_formatting.add(pr, CellIsRule(operator="equal", formula=['"OK"'], font=f(9, True, "1E8449"), fill=fill(GREEN_SOFT)))
dash.conditional_formatting.add(pr, CellIsRule(operator="equal", formula=['"REORDENAR"'], font=f(9, True, "B9770E"), fill=fill("FEF5E7")))
dash.conditional_formatting.add(pr, CellIsRule(operator="equal", formula=['"AGOTADO"'], font=f(9, True, RED), fill=fill(RED_SOFT)))
dash.conditional_formatting.add(f"F29:F{28 + NPROD_DASH}", CellIsRule(operator="equal", formula=['"Definir"'], font=f(9, True, RED), fill=fill("FFD6D6")))
dash.conditional_formatting.add(f"I29:I{28 + NPROD_DASH}", DataBarRule(start_type="num", start_value=0, end_type="max", color="8E44AD"))
dash.cell(tr + 1, 2, "Muestra los primeros 8 productos del catálogo. Detalle completo en la hoja INVENTARIO.").font = f(8, color="888888", italic=True)

# Gastos por categoría (tabla para gráfico)
GR0 = tr + 3
section(dash, GR0, 2, 7, "GASTOS POR CATEGORÍA (acumulado)", RED)
for i, cat in enumerate(CATS):
    rr = GR0 + 1 + i
    dash.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=5)
    dash.cell(rr, 2, cat).font = f(9)
    dash.cell(rr, 6, f'=SUMIFS(GASTOS!$J${G_FIRST}:$J${G_LAST},GASTOS!$C${G_FIRST}:$C${G_LAST},B{rr})').number_format = CUR
    dash.cell(rr, 6).font = f(9, True)
    dash.cell(rr, 7, f"=IF(SUM($F${GR0 + 1}:$F${GR0 + len(CATS)})>0,F{rr}/SUM($F${GR0 + 1}:$F${GR0 + len(CATS)}),0)").number_format = PCT
    dash.cell(rr, 7).font = f(9, color="666666")
    for cc in range(2, 8):
        dash.cell(rr, cc).fill = fill(LIGHT if i % 2 == 0 else WHITE)
        dash.cell(rr, cc).border = Border(bottom=Side(style="hair", color=MID))
dash.conditional_formatting.add(f"F{GR0 + 1}:F{GR0 + len(CATS)}", DataBarRule(start_type="num", start_value=0, end_type="max", color=RED))
rr = GR0 + 1 + len(CATS)
dash.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=5)
dash.cell(rr, 2, "TOTAL GASTOS")
dash.cell(rr, 6, f"=SUM(F{GR0 + 1}:F{rr - 1})").number_format = CUR
for cc in range(2, 8):
    dash.cell(rr, cc).fill = fill(NAVY); dash.cell(rr, cc).font = f(9, True, WHITE)

# Mensual (tabla para gráficos)
section(dash, GR0, 9, 13, "RESUMEN MENSUAL DEL AÑO", NAVY2)
header_row(dash, GR0 + 1, ["Mes", "Ventas brutas", "Ingreso neto", "Utilidad neta"], col0=9, color=TEAL, height=22)
dash.merge_cells(start_row=GR0 + 1, start_column=12, end_row=GR0 + 1, end_column=13)
for m in range(12):
    rr = GR0 + 2 + m
    col = CL(3 + m)
    dash.cell(rr, 9, MESES[m])
    dash.cell(rr, 10, f"=ESTADO_RESULTADOS!{col}{ER_VB}").number_format = CUR
    dash.cell(rr, 11, f"=ESTADO_RESULTADOS!{col}{ER_IN}").number_format = CUR
    dash.cell(rr, 12, f"=ESTADO_RESULTADOS!{col}{ER_UN}").number_format = CUR
    dash.merge_cells(start_row=rr, start_column=12, end_row=rr, end_column=13)
    for cc in range(9, 14):
        dash.cell(rr, cc).font = f(9, cc == 9)
        dash.cell(rr, cc).fill = fill(LIGHT if m % 2 == 0 else WHITE)
        dash.cell(rr, cc).border = Border(bottom=Side(style="hair", color=MID))
        dash.cell(rr, cc).alignment = Alignment(horizontal="center" if cc == 9 else "right")

# ---------- Gráficos ----------
CH0 = GR0 + 16
section(dash, CH0, 2, 13, "ANÁLISIS GRÁFICO", NAVY2)

bar = BarChart()
bar.type = "col"
bar.title = "Ventas vs. Utilidad neta por mes"
bar.style = 10
bar.y_axis.title = "USD"
bar.y_axis.numFmt = "$#,##0"
bar.y_axis.majorGridlines = None
data = Reference(dash, min_col=10, max_col=11, min_row=GR0 + 1, max_row=GR0 + 13)
bar.add_data(data, titles_from_data=True)
bar.set_categories(Reference(dash, min_col=9, min_row=GR0 + 2, max_row=GR0 + 13))
bar.series[0].graphicalProperties.solidFill = NAVY
bar.series[1].graphicalProperties.solidFill = TEAL
line = LineChart()
line.add_data(Reference(dash, min_col=12, min_row=GR0 + 1, max_row=GR0 + 13), titles_from_data=True)
line.series[0].graphicalProperties.line.solidFill = GOLD
line.series[0].graphicalProperties.line.width = 32000
line.series[0].smooth = True
bar += line
bar.legend.position = "b"
bar.height, bar.width = 8.5, 17
dash.add_chart(bar, f"B{CH0 + 1}")

dn = DoughnutChart()
dn.title = "Composición de activos"
dn.add_data(Reference(dash, min_col=11, min_row=20, max_row=23), titles_from_data=False)
dn.set_categories(Reference(dash, min_col=9, min_row=20, max_row=23))
dn.holeSize = 55
from openpyxl.chart.series import DataPoint
for i, colr in enumerate([TEAL, "5DADE2", "0071CE", "8E44AD"]):
    pt = DataPoint(idx=i)
    pt.graphicalProperties.solidFill = colr
    dn.series[0].dPt.append(pt)
dn.dataLabels = DataLabelList()
dn.dataLabels.showPercent = True
for k in ('showVal','showCatName','showSerName','showLegendKey'):
    setattr(dn.dataLabels, k, False)
dn.legend.position = "r"
dn.height, dn.width = 8.5, 12.5
dash.add_chart(dn, f"I{CH0 + 1}")

b2 = BarChart()
b2.type = "bar"
b2.title = "Valor del inventario por producto"
b2.style = 10
b2.add_data(Reference(dash, min_col=9, min_row=28, max_row=28 + NPROD_DASH), titles_from_data=True)
b2.set_categories(Reference(dash, min_col=2, min_row=29, max_row=28 + NPROD_DASH))
b2.series[0].graphicalProperties.solidFill = "8E44AD"
b2.x_axis.scaling.orientation = "maxMin"
b2.y_axis.numFmt = "$#,##0"
b2.y_axis.majorGridlines = None
b2.legend = None
b2.height, b2.width = 8.5, 17
dash.add_chart(b2, f"B{CH0 + 19}")

b3 = BarChart()
b3.type = "bar"
b3.title = "Gastos por categoría"
b3.style = 10
b3.add_data(Reference(dash, min_col=6, min_row=GR0 + 1, max_row=GR0 + len(CATS)), titles_from_data=False)
b3.set_categories(Reference(dash, min_col=2, min_row=GR0 + 1, max_row=GR0 + len(CATS)))
b3.series[0].graphicalProperties.solidFill = RED
b3.x_axis.scaling.orientation = "maxMin"
b3.y_axis.numFmt = "$#,##0"
b3.y_axis.majorGridlines = None
b3.legend = None
b3.height, b3.width = 8.5, 12.5
dash.add_chart(b3, f"I{CH0 + 19}")

dash.freeze_panes = "A4"

# =====================================================================
# GUIA
# =====================================================================
banner(guia, "COTIMIX · GUÍA DE USO DEL SISTEMA", "Cómo registrar la operación diaria y leer los reportes.", 6)
guia.column_dimensions["B"].width = 26
guia.column_dimensions["C"].width = 92
blocks = [
    ("CÓDIGO DE COLORES", None),
    ("Celda amarilla, texto azul", "Dato que USTED captura o puede modificar (inputs)."),
    ("Celda blanca, texto negro", "Fórmula automática. No la sobrescriba."),
    ("Texto verde", "Vínculo a otra hoja (por ejemplo, el capital en BANCO viene de CONFIG)."),
    ("Rojo claro", "Alerta: dato pendiente (precio sin definir) o resultado negativo."),
    ("FLUJO DE TRABAJO", None),
    ("1. CONFIG", "Revise parámetros: capital, reserva, recargo Citi (2.9%), comisión Walmart y margen objetivo."),
    ("2. INVENTARIO", "Defina el PRECIO DE VENTA de cada producto (columna N). Ajuste la comisión Walmart de su categoría "
                      "y el costo de envío por unidad. Para un producto nuevo, agregue un ID (P003, P004...) en una fila vacía."),
    ("3. COMPRAS", "Registre cada compra de mercancía. Si elige 'Tarjeta Citi', el 2.9% se suma al costo real automáticamente "
                   "y la deuda Citi aumenta. El costo unitario del inventario es el promedio ponderado del costo real."),
    ("4. VENTAS", "Una fila por orden/producto: fecha, # orden, ID y cantidad. El precio es opcional (si lo deja vacío usa el del catálogo)."),
    ("5. GASTOS", "Publicidad, envíos, empaques, etc. Use 'Descontado por Walmart' para cargos que Walmart resta de la liquidación."),
    ("6. WALMART", "Registre los depósitos que Walmart envía a su banco y los movimientos de la reserva (retenciones y liberaciones)."),
    ("7. BANCO", "Registre aportes adicionales, retiros del propietario, pagos a la tarjeta Citi y otros ingresos."),
    ("8. Reportes", "DASHBOARD, ESTADO_RESULTADOS y BALANCE se actualizan solos. Verifique que el BALANCE diga '✔ CUADRA'."),
    ("EJEMPLO DE CAPTURA (VENTAS)", None),
    ("Formato de una fila", "Fecha: 05/11/2026 · # Orden: 1089456123 · ID: P002 · Cantidad: 2 · Precio unitario: 79.99 · "
                            "Envío/otros cargos: 0  →  Venta bruta $159.98, comisión 15% $24.00, costo real $86.44, utilidad $49.54."),
    ("SUPUESTOS DOCUMENTADOS", None),
    ("Reserva Walmart", "$1,190.60 es dinero de COTIMIX retenido por Walmart: se muestra como activo RESTRINGIDO, separado del efectivo."),
    ("Origen de la reserva", "Se asume que proviene de ventas previas a la apertura → 'Resultados acumulados de apertura' en el patrimonio. Editable en CONFIG."),
    ("Deuda Citi", "Se asume que la compra de mouse ($1,728.72) aún no se ha pagado: aparece como pasivo hasta registrar el pago en BANCO."),
    ("Efectivo inicial", "Capital $5,000 − cafeteras pagadas en efectivo $2,070 = $2,930 disponibles a la apertura."),
    ("Comisión Walmart", "15% por defecto (tarifa referencial estándar). Verifique la de cada categoría en Seller Center."),
]
r = 5
for a, b in blocks:
    if b is None:
        r += 1
        section(guia, r, 2, 3, a, TEAL)
    else:
        guia.cell(r, 2, a).font = f(10, True, NAVY)
        guia.cell(r, 3, b).font = f(10)
        guia.cell(r, 3).alignment = Alignment(wrap_text=True, vertical="center")
        guia.cell(r, 2).alignment = Alignment(vertical="center")
        guia.row_dimensions[r].height = 30 if len(b) > 95 else 18
        for cc in (2, 3):
            guia.cell(r, cc).border = Border(bottom=Side(style="hair", color=MID))
    r += 1
guia["B7"].fill = fill(INPUT_FILL); guia["B7"].font = f(10, True, "0000FF")
guia["B9"].font = f(10, True, "008000")
guia["B10"].fill = fill("FFD6D6")

# ---------- Ajustes finales ----------
for ws in wb.worksheets:
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.oddFooter.center.text = "COTIMIX · Sistema de control financiero · Página &P de &N"
    ws.oddFooter.center.size = 8
wb.calculation.fullCalcOnLoad = True
wb.active = 0
wb.save(OUT)
print("saved", OUT)
